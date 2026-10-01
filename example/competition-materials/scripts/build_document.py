"""Build the competition explanation from reviewed local sources and clean screenshots."""
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / '学途智伴技能说明文档.docx'


def font(style, size, bold=False):
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
    for name in ('ascii', 'hAnsi', 'eastAsia'):
        fonts.set(qn('w:' + name), 'Microsoft YaHei')
    for name in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme'):
        fonts.attrib.pop(qn('w:' + name), None)
    color = style._element.rPr.find(qn('w:color'))
    if color is not None:
        color.attrib.pop(qn('w:themeColor'), None)


def body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead:
        p.add_run(bold_lead).bold = True
    p.add_run(text)
    return p


def heading(doc, text, level=1):
    return doc.add_heading(text, level)


def page_section(doc, text):
    p = heading(doc, text)
    p.paragraph_format.page_break_before = True
    return p


def picture(doc, file, caption, width=6.65):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    inline = p.add_run().add_picture(str(ROOT / 'assets' / file), width=Inches(width))
    inline._inline.docPr.set('descr', caption)
    c = doc.add_paragraph(caption, style='Caption')
    c.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return c


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for i, width in enumerate(widths):
        t.columns[i].width = Inches(width)
    for n, values in enumerate([headers] + rows):
        row = t.rows[0] if n == 0 else t.add_row()
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if n == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for i, value in enumerate(values):
            cell = row.cells[i]
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = value
            pr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), '173B55' if n == 0 else ('F2F6FA' if n % 2 == 0 else 'FFFFFF'))
            pr.append(shade)
            borders = OxmlElement('w:tcBorders')
            for edge in ('top', 'left', 'bottom', 'right'):
                e = OxmlElement('w:' + edge)
                for name, val in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                    e.set(qn('w:' + name), val)
                borders.append(e)
            pr.append(borders)
            margins = OxmlElement('w:tcMar')
            for edge, size in [('top', 85), ('bottom', 85), ('left', 110), ('right', 110)]:
                e = OxmlElement('w:' + edge)
                e.set(qn('w:w'), str(size)); e.set(qn('w:type'), 'dxa'); margins.append(e)
            pr.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.18
                for r in p.runs:
                    r.font.size = Pt(10)
                    r.font.bold = n == 0
                    r.font.color.rgb = RGBColor(255, 255, 255) if n == 0 else RGBColor(0, 0, 0)
    gap = doc.add_paragraph()
    gap.paragraph_format.space_after = Pt(2)
    gap.paragraph_format.space_before = Pt(0)
    gap.paragraph_format.line_spacing = 0.4
    return t


def link(doc, label, url):
    p = doc.add_paragraph(style='Source')
    h = OxmlElement('w:hyperlink')
    rel = p.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    h.set(qn('r:id'), rel)
    r = OxmlElement('w:r'); pr = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '173B55'); pr.append(color)
    r.append(pr); text = OxmlElement('w:t'); text.text = label; r.append(text); h.append(r); p._p.append(h)
    return p


def step(doc, number, title, text):
    return body(doc, text, f'{number} {title}  ')


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = Inches(.7); section.bottom_margin = Inches(.65)
    section.left_margin = Inches(.78); section.right_margin = Inches(.78)
    section.footer_distance = Inches(.3)
    for name, size, bold in [('Normal', 11, False), ('Title', 22, True), ('Subtitle', 12, False), ('Heading 1', 15, True), ('Heading 2', 11.5, True), ('Caption', 9, False)]:
        font(doc.styles[name], size, bold)
    normal = doc.styles['Normal'].paragraph_format
    normal.line_spacing = 1.25; normal.space_after = Pt(7)
    for name in ('Heading 1', 'Heading 2'):
        p = doc.styles[name].paragraph_format
        p.space_before = Pt(11); p.space_after = Pt(6); p.keep_with_next = True
    doc.styles['Title'].paragraph_format.space_after = Pt(5)
    doc.styles['Subtitle'].paragraph_format.space_after = Pt(9)
    doc.styles['Caption'].paragraph_format.space_after = Pt(7)
    doc.styles['Caption'].paragraph_format.line_spacing = 1.15
    source = doc.styles.add_style('Source', 1)
    source.base_style = doc.styles['Normal']; font(source, 9)
    source.paragraph_format.space_after = Pt(3)
    # Remove title-rule residue and all decorative paragraph borders.
    for style in doc.styles:
        ppr = style._element.find(qn('w:pPr'))
        if ppr is not None:
            for border in list(ppr.findall(qn('w:pBdr'))):
                ppr.remove(border)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)

    doc.add_paragraph('学途智伴技能说明文档', style='Title')
    doc.add_paragraph('大学生长期自适应学习智能体', style='Subtitle')
    body(doc, '队伍  重邮FFBond    赛道  校园生活    Skill  Teach Pro 1.1.0-rc.4')
    heading(doc, '作品简介')
    body(doc, '学途智伴面向大学生自主学习、跨专业入门和竞赛备赛。我们将学习目标、基础评估、逐课教程与本地作答连成持续的教学过程：先了解起点，每次交付一节完整课，再依据解释、练习和疑难调整下一步。')
    body(doc, '作品名称为“学途智伴——大学生长期自适应学习智能体”。教学能力以 teach-pro Skill 提供，可在支持技能调用与工作区读写的 Agent 环境中使用。学生通过浏览器学习，在同一课程目录保留长期记录。')
    body(doc, '比赛 Demo 选用“智鉴 Agent：智能体开发与安全实战”。学习目标是开发安全情报 Agent，起点设定为基本 Python 与 API 经验。示例展示基础课、信任边界补救课和工具调用推进课，附入门评估、一手阅读中心与统一课程设置。')
    picture(doc, 'course-home.png', '图 1 智鉴 Agent 展示课程首页  课程页面经人工校核')
    body(doc, '课程采用主题相关的轻量网格背景、可折叠侧栏和统一按钮样式。课内用原创图解解释机制，阅读资源指向具体章节或论文片段，并给出可检查的短任务。')

    page_section(doc, '设计思路')
    heading(doc, '根据证据选择教学任务', 2)
    body(doc, '我们把长期路线与当前教学决策分开。路线只规划候选能力阶段，实际续课读取最新作答：解释不足时补充理解任务，出现具体误区时换表达与情境进行补救，前置内容得到支持后再进入新目标。页面浏览和一句“继续”只能触发检查，不能单独证明掌握。')
    heading(doc, '用文件维持学习连续性', 2)
    body(doc, '课程目录保存学习使命、画像、路线、教学记录，以及按课组织的答案和疑难。学生只需填写页面，再回到同一工作区请求继续。AI 读取这些文件，先回答疑问，再选择一节课并更新状态；生成长课程时仍能回到有来源的学习记录。')
    picture(doc, 'architecture.png', '图 2 教学与文件连续性  可选 Tutor 与课程生成 AI 分工独立')
    heading(doc, '逐课交付与精确导学', 2)
    body(doc, '每节课围绕当前目标组织图解、示例、练习和反馈，记录解释、应用与迁移表现。外部资源标明“看哪里、为什么看、看后做什么”，帮助学生带着问题阅读，而不是面对整套资料自行筛选。支持图像、内嵌视频和字幕；本 Demo 使用原创图解及短阅读任务。')
    body(doc, '这套设计的重点是让下一节课有可追溯的输入、教学理由与产出。课内问答提供即时帮助，课程生成 AI 则负责持续安排教学，两者通过本地文件衔接。')

    page_section(doc, '技术实现与数据边界')
    body(doc, 'SKILL.md 组织教学行为，配套格式规范、HTML 模板和复用资源。网页负责学习交互；Markdown 表达课程状态；本地 Python 服务将作答写入 JSON，并提供可选答疑接口。课程服务使用 Python 标准库，无需第三方 Python 包。')
    table(doc, ['文件或组件', '职责'], [
        ['SKILL.md 与配套规范', '约束评估、逐课生成、证据判断和内容校核'],
        ['MISSION 与 LEARNER-PROFILE', '保存目标、偏好和有来源的能力判断'],
        ['COURSE 与 ROADMAP', '区分已交付课程、当前状态和下一候选'],
        ['lessons 与 practice', '展示课文、入门评估和练习输入'],
        ['learner-submissions', '按课保存原始作答与疑难 JSON'],
        ['learner-chats', '按课保存可选 Tutor 的聊天原文'],
        ['serve_course 与共享资源', '提供回环服务、自动同步、导航和设置交互'],
    ], [2.35, 4.59])
    heading(doc, '本地保存与访问范围', 2)
    body(doc, '服务只监听 127.0.0.1，并限制静态文件与提交入口。直接打开 HTML 可阅读并在浏览器保存；使用启动器后，答案才会写入课程目录，供具有该目录访问权限的 AI 读取。读取发生在用户请求续课时，而非后台自动开课。')
    heading(doc, '密钥与模型上下文', 2)
    body(doc, '课程设置保存非密钥连接参数；API Key 仅保留在本次服务进程内存中，重启后重新输入。获取模型列表和连接测试不发送课程正文。学员点击发送后，答疑结合本课正文与图注、当前问题、选中片段和最近五个完整成功回合，不自动读取其他课程或完整学员画像。')
    body(doc, '本地保存不等于本地推理，云端答疑会把限定上下文发往所选服务。作答、聊天和配置已加入忽略规则；公开 Demo 不包含原始学员提交、聊天或密钥。本服务用于单用户本机学习。')

    page_section(doc, '使用说明')
    step(doc, 1, '导入并显式调用', '将 teach-pro 目录导入支持 Skill 和工作区读写的 Agent。TeleAgent 中可输入下面的示例，先评估再开课。')
    body(doc, '@teach-pro 我想学习 Agent 开发与 Agent 安全，会基本 Python 和模型 API，每周约五小时，偏好图解，目标是做安全情报 Agent。请先评估我的基础再开课。')
    step(doc, 2, '完成起点评估', '在新课程的评估页面填写选择题和开放题。可以说明不记得，并记录是否参考过提示。首次评估之后，请 AI 读取答案并生成第一课。')
    step(doc, 3, '启动课程并学习', '静态阅读打开 index.html。自动写入本地文件需要 Python 3.11+：Windows 双击 start-course.cmd，macOS 双击 start-course.command，Linux 运行 sh start-course.sh；也可运行 python serve_course.py。')
    picture(doc, 'assessment.png', '图 3 入门评估的选择题与检查按钮  四道选择题和两道开放题均可保存', width=6.4)
    step(doc, 4, '保存答案并请求续课', '完成当前课的解释、练习和疑难，确认页面显示“已写入课程目录”。回到同一工作区，请 AI 读取最新作答，先回答疑问，再决定下一节课。断线时可下载 JSON 备份。')
    body(doc, '运行环境的 Skill 导入规则可能不同。TeleAgent 曾对 .cmd 扩展名提示限制；本地 Demo 的启动器可独立使用。示例中的三节课是逐次生成后的校核整理，新课程仍按实际学习证据逐节交付。')

    page_section(doc, '课程设置与课内答疑')
    body(doc, 'AI 答疑按需启用。所有课节共用左侧置顶“课程设置”，课程页只提供当前课的问答入口。这里的连接仅用于课内 Tutor，不改变 TeleAgent 等环境中课程生成 AI 的模型。')
    heading(doc, '以 DeepSeek 为例', 2)
    step(doc, 1, '选择服务并输入 Key', '从服务类型选择 DeepSeek，默认地址自动填为 https://api.deepseek.com。输入自己的 API Key；默认以圆点遮蔽，小眼睛可切换显示。')
    step(doc, 2, '获取并选择模型', '点击“获取模型”，从实际列表选择 deepseek-flash，再保存并启用。不支持列表时可手动填写模型 ID。地址、输出预算与思考模式放在高级设置中。')
    picture(doc, 'deepseek-settings.png', '图 4 简化后的 DeepSeek 设置  图中密钥为无效演示值', width=5.7)
    step(doc, 3, '保存并测试', '保存后 Key 输入框清空，密钥留在服务内存中。测试连接发送简短请求；普通 Flash 答疑默认关闭思考，聊天预算默认 4096 token，连接测试预算独立。')
    step(doc, 4, '返回当前课提问', '点击“问 AI”进入大弹窗，直接提问或继续追问，无需粘贴课文或重复勾选。未配置也能继续学习。')

    page_section(doc, '课内对话体验')
    body(doc, '答疑窗口把注意力集中在当前疑问上：顶部显示课节与模型，正文以连续文档流呈现，输入区固定在底部。标题、列表、强调、代码和表格按 Markdown 排版，代码支持复制，长内容在局部滚动。')
    picture(doc, 'tutor-dialog.png', '图 5 DeepSeek Flash 本次真实第三轮答疑  模型回答原样展示', width=6.4)
    body(doc, '学员可以先给出自己的判断，再让模型指出推理缺口。演示中，三轮交流依次讨论学校公告、新邮箱与外发权限，最后澄清系统提示与运行时授权的区别。即时答疑辅助理解，学员的解释和应用仍须另行检查。')
    heading(doc, '保持对话连续', 2)
    body(doc, '聊天原始 Markdown 自动保存到本课 JSON。关闭窗口保留草稿，已发送请求继续接收；刷新恢复历史，阅读旧消息时新回答不会强制跳到底。支持引用课文、Enter 发送、Shift 加 Enter 换行、输入法保护，以及备份和确认清空。')
    heading(doc, '请求状态清晰', 2)
    body(doc, '服务区分正常回答、部分正文、预算耗尽、异常结构和连接失败。未完成回答明确标记，学员可主动重试最近问题；重试关联原记录，不重复插入提问。有限诊断记录结束原因、用量与耗时，不保存思考原文或认证信息。')

    page_section(doc, '演示内容与验证结果')
    body(doc, 'Demo 共八页，包含三节校核课程、阅读中心与决策回放。新版画面依次展示评估、图解、疑问、配置、真实答疑、保存与续课，时长不足三分钟，密钥始终遮蔽。配音另行制作。')
    heading(doc, '教学决策观察', 2)
    table(doc, ['合成输入', '观察到的教学任务'], [
        ['关键开放题内容不足', '交付巩固课，补充解释与迁移证据'],
        ['混淆公告外观与外发授权', '先回答疑难，再生成信任边界补救课'],
        ['新内容能区分身份与动作权限', '处理数据流连接疑问，生成工具调用推进课'],
    ], [2.65, 4.29])
    body(doc, '三种选择均在 TeleAgent v2.5.2 中观察到，使用合成即时作答，推进分支基于人工修订的前课。原生课文仍有概念与状态错误，展示课程已校核；无人干预闭环和长期学习效果尚未验证。')
    heading(doc, '页面与文件链路验收', 2)
    table(doc, ['检查对象', '结果与范围'], [
        ['单元与结构检查', 'Node 16/16、Python 37/37；八页结构及 Skill 校验通过'],
        ['弹窗与排版', '1440、1280、390px 布局通过；Markdown、输入法、滚动与关闭恢复正常'],
        ['入门答案与课末疑难', '六字段落盘、文件恢复、导出与清空通过；课间隔离正常'],
        ['真实模型请求', 'DeepSeek Flash 模型列表、连接测试、三轮答疑均成功'],
        ['发布权限', '.sh 与 .command 的 Git 和 ZIP 模式均为 755'],
    ], [2.05, 4.89])
    body(doc, '页面回归使用合成接口，真实 DeepSeek 请求独立记录。当前验收限于 Windows，未验证 macOS/Linux 实机启动。教学分支保留原生失败项，不因界面升级改记为通过。')
    base = 'https://github.com/RoamerFly/teach-pro-skill/blob/main/'
    link(doc, '项目仓库与使用入口', 'https://github.com/RoamerFly/teach-pro-skill')
    link(doc, 'Tutor 体验与真实演示验收  2026 年 10 月 2 日', base + 'evals/records/2026-10-02-tutor-materials-review.md')
    link(doc, 'G3 原生续课与人工修订记录', base + 'evals/records/2026-10-01-test5-g3-recovery-review.md')
    props = doc.core_properties
    props.title = '学途智伴技能说明文档'; props.author = '重邮FFBond'
    props.subject = '校园生活赛道作品简介设计思路与使用说明'
    props.keywords = 'Teach Pro, 1.1.0-rc.4, 自适应教学, 本地学习记录'
    props.comments = ''; props.last_modified_by = '重邮FFBond'
    props.created = datetime(2026, 10, 1, tzinfo=timezone.utc)
    props.modified = datetime(2026, 10, 2, tzinfo=timezone.utc)
    doc.save(OUT)
    print(OUT)


if __name__ == '__main__':
    build()
