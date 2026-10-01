"""Read-only structural checks; visual review is still required for delivery."""
import hashlib
from pathlib import Path
import re
from zipfile import ZipFile
from docx import Document
from docx.oxml.ns import qn

file = Path(__file__).resolve().parents[1] / '学途智伴技能说明文档.docx'
doc = Document(file)
expected = ['作品简介', '设计思路', '技术实现与数据边界', '使用说明', '课程设置与课内答疑', '演示内容与验证结果']
assert [p.text for p in doc.paragraphs if p.style.name == 'Heading 1'] == expected
assert len(doc.inline_shapes) == 4
assert len(doc.tables) == 3
assert doc.core_properties.author == '重邮FFBond'
assert doc.core_properties.last_modified_by == '重邮FFBond'
assert abs(doc.sections[0].page_width / 914400 - 8.5) < .001
assert abs(doc.sections[0].page_height / 914400 - 11) < .001
for shape in doc.inline_shapes:
    assert shape._inline.docPr.get('descr')
for name in ('Title', 'Subtitle', 'Heading 1', 'Heading 2'):
    assert str(doc.styles[name].font.color.rgb) == '000000', name
    assert not doc.styles[name]._element.findall('.//' + qn('w:pBdr'))
for table in doc.tables:
    assert table.rows[0]._tr.get_or_add_trPr().find(qn('w:tblHeader')) is not None
    for row in table.rows:
        for cell in row.cells:
            borders = cell._tc.get_or_add_tcPr().find(qn('w:tcBorders'))
            assert borders is not None
            assert all(edge.get(qn('w:color')) == 'D9D9D9' for edge in borders)
with ZipFile(file) as z:
    assert z.testzip() is None
    xml = '\n'.join(z.read(name).decode('utf-8') for name in z.namelist() if name.endswith('.xml') or name.endswith('.rels'))
    assert not re.search(r'sk-[A-Za-z0-9]{16,}', xml)
    assert 'learner-submissions/' not in '\n'.join(z.namelist())
    assert 'learner-chats/' not in '\n'.join(z.namelist())
    assert 'w:del ' not in xml and 'w:ins ' not in xml
    assert not any(name.startswith('word/comments') for name in z.namelist())
    hyperlinks = [rel for rel in doc.part.rels.values() if rel.reltype.endswith('/hyperlink')]
    assert len(hyperlinks) == 3
    assert all(rel.target_ref.startswith('https://github.com/RoamerFly/teach-pro-skill') for rel in hyperlinks)
print('PASS: six sections, four annotated images, three tables and three public source links')
print('SHA256:', hashlib.sha256(file.read_bytes()).hexdigest().upper())
