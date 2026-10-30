"""Structural audit only; every rendered page must also be reviewed visually."""
from pathlib import Path
from zipfile import ZipFile
import hashlib
import json
import re
from docx import Document
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[1]
file = root / '学途智伴技能说明文档.docx'
doc = Document(file)
assert [p.text for p in doc.paragraphs if p.style.name == 'Heading 1'] == ['作品简介', '设计思路', '使用说明', '课程设置与课内答疑', '验证结果与参赛演示']
assert len(doc.inline_shapes) == 5
assert len(doc.tables) == 1
assert doc.core_properties.author == doc.core_properties.last_modified_by == '重邮FFBond'
assert '1.1.0-rc.18' in doc.core_properties.keywords
assert '上下文设置' in '\n'.join(p.text for p in doc.paragraphs)
assert '每道练习后直接作答' in '\n'.join(p.text for p in doc.paragraphs)
for name in ('Title', 'Subtitle', 'Heading 1', 'Heading 2'):
    assert str(doc.styles[name].font.color.rgb) == '000000'
    assert not doc.styles[name]._element.findall('.//' + qn('w:pBdr'))
for shape in doc.inline_shapes:
    assert shape._inline.docPr.get('descr')
for table in doc.tables:
    assert table.rows[0]._tr.get_or_add_trPr().find(qn('w:tblHeader')) is not None
    for row in table.rows:
        assert not row._tr.findall('.//' + qn('w:trHeight'))
        for cell in row.cells:
            borders = cell._tc.get_or_add_tcPr().find(qn('w:tcBorders'))
            assert borders is not None
            assert all(edge.get(qn('w:color')) == 'D9D9D9' for edge in borders)
with ZipFile(file) as archive:
    assert archive.testzip() is None
    xml = '\n'.join(archive.read(n).decode('utf8') for n in archive.namelist() if n.endswith('.xml'))
    assert not re.search(r'sk-[A-Za-z0-9]{16,}', xml)
    assert 'w:del ' not in xml and 'w:ins ' not in xml
    assert not any(n.startswith('word/comments') for n in archive.namelist())
    figure = (root / 'assets/teach-pro-workflow-architecture-v2-connected.png').read_bytes()
    assert any(archive.read(n) == figure for n in archive.namelist() if n.startswith('word/media/'))
    duration = round(json.loads((root / 'dialogue-video-audit.json').read_text(encoding='utf8'))['seconds'])
    assert f'{duration // 60} 分 {duration % 60} 秒' in xml
    assert '工作流程与系统架构' in xml and '教学 AI' in xml and '课内 Tutor' in xml
print('PASS: five sections, five annotated images, one matrix, rc.18 metadata, context settings, no secret patterns')
print('SHA256:', hashlib.sha256(file.read_bytes()).hexdigest())
